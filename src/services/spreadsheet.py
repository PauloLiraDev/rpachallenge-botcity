"""Leitura da planilha challenge.xlsx."""

from collections.abc import Iterator
from pathlib import Path

from openpyxl import load_workbook
from pydantic import ValidationError

from src.domain.models import Person
from src.exceptions import InvalidInputError

HEADER_TO_FIELD: dict[str, str] = {
    "first name": "first_name",
    "last name": "last_name",
    "company name": "company_name",
    "role in company": "role",
    "address": "address",
    "email": "email",
    "phone number": "phone",
}


def _normalize_header(value: object) -> str:
    return " ".join(str(value or "").split()).lower()


def _map_columns(header: tuple[object, ...]) -> dict[int, str]:
    columns = {
        index: HEADER_TO_FIELD[name]
        for index, cell in enumerate(header)
        if (name := _normalize_header(cell)) in HEADER_TO_FIELD
    }
    missing = set(HEADER_TO_FIELD.values()) - set(columns.values())
    if missing:
        raise InvalidInputError(f"Colunas ausentes na planilha: {sorted(missing)}")
    return columns


def _rows(path: Path) -> Iterator[tuple[object, ...]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        yield from workbook.active.iter_rows(values_only=True)
    finally:
        workbook.close()


def read_people(path: Path) -> list[Person]:
    """Lê a planilha e devolve uma pessoa por linha preenchida."""
    rows = _rows(path)
    columns = _map_columns(next(rows, ()))
    people: list[Person] = []
    for line, row in enumerate(rows, start=2):
        if all(cell in (None, "") for cell in row):
            continue
        values = {field: row[index] for index, field in columns.items()}
        try:
            people.append(Person.model_validate(values))
        except ValidationError as exc:
            raise InvalidInputError(f"Linha {line} inválida: {exc.error_count()} erro(s)") from exc
    return people
